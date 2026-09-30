# -*- coding: utf-8 -*-
"""把投稿信 Markdown 渲染成 .docx（投稿信不受 10pt/1.5 的正文版式约束，用常规商务版式）。

★ 投稿信是 PR 的**必交件**，且官方要求在里面回答三个问题（L801–L810）。本脚本只负责排版，
  内容由 `09_CoverLetter_PR.md` 单一来源提供 —— 不允许"docx 里另写一份"。
"""


# ── 复现包统一根：`_repro_root.py`（与本文件同目录）──────────────────────────────
# RP(*parts) = 作者树相对路径 -> 绝对路径（作者树上原样；放行树上查前缀映射表）；
# NR(*parts) = **未随包发布**的作者侧路径（放行树上落到 _NOT_RELEASED/，使失败可见）。
try:
    from _repro_root import resolve as RP, not_released as NR
except ImportError:                      # 只拷走单个脚本时：就地反推仓库根，无前缀映射表
    import os as _o
    _r = _o.environ.get('PAPERB_ROOT') or _o.path.dirname(_o.path.dirname(_o.path.abspath(__file__)))
    RP = lambda *p: _o.path.join(_r, *p)
    NR = lambda *p: _o.path.join(_r, '_NOT_RELEASED', *p)
import io
import os
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, RP('analysis', 'work'))
import build_pr_docx as B          # 复用同一套行内解析与页面设置
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

PKG = NR('review_pkg_20260919')
SRC = NR('review_pkg_20260919', '09_CoverLetter_PR.md')
OUT = NR('review_pkg_20260919', '09_CoverLetter_PR.docx')

md = io.open(SRC, encoding='utf-8', newline='\n').read()
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2.54)
sec.left_margin = sec.right_margin = Cm(2.54)
st = doc.styles['Normal']
st.font.name = 'Times New Roman'
st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.15

lines = md.split('\n')
i = 0
while i < len(lines):
    ln = lines[i]
    s = ln.strip()
    if not s:
        i += 1
        continue
    if s.startswith('> '):                       # 引用块
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        B.add_inline(p, s[2:], size=11)
        i += 1
        continue
    if re.match(r'^#{1,4}\s', s):                # 标题
        lvl = len(re.match(r'^(#+)', s).group(1))
        p = doc.add_paragraph()
        B.add_inline(p, re.sub(r'^#+\s*', '', s), size=(13 if lvl <= 2 else 11.5), bold_all=True)
        p.paragraph_format.space_before = Pt(8)
        i += 1
        continue
    if s.startswith('|'):                        # 表格
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not re.match(r'^[\s\-:|]+$', lines[i].strip().strip('|')):
                rows.append(cells)
            i += 1
        ncol = max(len(r) for r in rows)
        t = doc.add_table(rows=0, cols=ncol)
        B.set_cell_borders(t)
        for ri, row in enumerate(rows):
            cs = t.add_row().cells
            for ci in range(ncol):
                par = cs[ci].paragraphs[0]
                par.paragraph_format.line_spacing = 1.0
                par.paragraph_format.space_after = Pt(0)
                B.add_inline(par, row[ci] if ci < len(row) else '', size=10.5, bold_all=(ri == 0))
        doc.add_paragraph()
        continue
    if re.match(r'^[-*]\s+', s):                 # 列表
        p = doc.add_paragraph(style='List Bullet')
        B.add_inline(p, re.sub(r'^[-*]\s+', '', s), size=11)
        i += 1
        continue
    if re.match(r'^\d+\.\s+', s):                # 有序列表
        p = doc.add_paragraph(style='List Number')
        B.add_inline(p, re.sub(r'^\d+\.\s+', '', s), size=11)
        i += 1
        continue
    if re.match(r'^-{3,}$', s):                  # 水平线 → 空段
        doc.add_paragraph()
        i += 1
        continue
    # 普通段落（连续行合并）
    buf = [s]
    i += 1
    while i < len(lines) and lines[i].strip() and not re.match(
            r'^(#{1,4}\s|\||-{3,}$|[-*]\s|\d+\.\s|>)', lines[i].strip()):
        buf.append(lines[i].strip())
        i += 1
    p = doc.add_paragraph()
    B.add_inline(p, ' '.join(buf), size=11)

doc.save(OUT)
print('输出 %s（%d B）' % (OUT, os.path.getsize(OUT)))

# 复核：三个问题都在，且"页面事实"与冻结的测量一致
txt = md
for need in ('1. Is the work compared with the state of the art?',
             '2. Which public datasets were used?',
             '3. Which validation metrics were used?'):
    assert need in txt, '投稿信缺少必答问题：%s' % need
print('三个必答问题齐全 ✓')
import json
m = json.loads(io.open(RP('measurement_pr_docx.json'), encoding='utf-8').read())
assert str(m['result']['pages']) in txt, '投稿信中的页数与实测不一致'
print('页数自述与实测一致（%d 页）✓' % m['result']['pages'])
